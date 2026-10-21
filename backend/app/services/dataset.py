"""结构化表格数据集：让 `{活动1名称}` 这类带序号的占位符能精确落到表格某一行。

为什么单独做这一层：
plan.md 里，网络安全接口协议是 excel，活动清单是「一行一个活动」的结构化数据。
如果只走全文检索，「活动1名称」这种占位符对不上任何字面内容，必然误判为缺失。
但只要认出「活动 / 序号 1 / 属性 名称」这个结构，就能直接定位到第 1 行的
「活动名称」列——准确率 100%，还能顺带把 Target Date 预填进结束时间。

所以策略是：结构化数据优先，命中即用；未命中再回落到文本检索。
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

SEQ_RE = re.compile(r"^([\u4e00-\u9fa5A-Za-z]{1,8}?)(\d{1,2})([\u4e00-\u9fa5A-Za-z]{1,12})$")

# 占位符属性 → 表格列名同义词。模板作者和协议编写者用词经常不一致，
# 这层映射就是用来吸收差异的（尤其 work product 有中英多种写法）。
ATTR_SYNONYMS: dict[str, list[str]] = {
    # 不要放「序号」——协议表里「序号」和「活动编号」常常并列，
    # 混用会让 {活动1编号} 错误地取到行号 1 而不是 CS-A01
    "编号": ["活动编号", "编号", "id", "no"],
    "名称": ["活动名称", "名称", "name", "事项"],
    "责任人": ["责任人", "责任方", "负责人", "执行方", "owner"],
    "交付物": ["work product", "工作产物", "工作成果", "交付物", "产物", "输出"],
    "开始时间": ["开始", "start", "起始"],
    "结束时间": ["target date", "结束", "截止", "提交时间", "计划完成", "end", "完成时间"],
    "时间": ["时间", "日期", "date"],
    "姓名": ["姓名", "name"],
    "部门": ["部门", "department"],
    "角色": ["角色", "岗位", "职责"],
    "联系方式": ["联系方式", "电话", "手机", "邮箱", "email"],
    "状态": ["状态", "status"],
}


@dataclass
class DataTable:
    source: str                 # 文件名
    name: str                   # sheet 名 / 表格标题
    headers: list[str]
    rows: list[list[str]] = field(default_factory=list)

    @property
    def header_blob(self) -> str:
        return " ".join(self.headers).lower()

    def col_index(self, attr: str) -> int | None:
        # 长候选词优先：「活动名称」应该先于「名称」命中，避免被更泛的列抢走
        candidates = sorted(ATTR_SYNONYMS.get(attr, [attr]), key=len, reverse=True)
        lowered = [h.lower() for h in self.headers]
        for key in candidates:
            for i, head in enumerate(lowered):
                if head == key:
                    return i
        for key in candidates:
            for i, head in enumerate(lowered):
                if key in head:
                    return i
        return None

    def index_of(self, header: str) -> int | None:
        """按列名精确取下标，供模型映射结果落地时校对用。"""
        target = (header or "").strip().lower()
        for i, head in enumerate(self.headers):
            if head.strip().lower() == target:
                return i
        return None

    def cell(self, row_index: int, col_index: int) -> str:
        if row_index < 0 or row_index >= len(self.rows):
            return ""
        row = self.rows[row_index]
        if col_index >= len(row):
            return ""
        return str(row[col_index]).strip()

    def data_rows(self) -> list[int]:
        """有效数据行下标。

        末尾的「合计 / 小计 / 备注」不是一条业务数据，扩行时不能把它算进去——
        多出一行内容为「合计」的活动，比少一行难看得多。
        """
        skip = ("合计", "总计", "小计", "备注", "说明")
        out: list[int] = []
        for i, row in enumerate(self.rows):
            head = " ".join(str(v).strip() for v in row[:2])
            if not head.strip():
                continue
            if any(k == head.strip() or head.strip().startswith(k) for k in skip):
                continue
            out.append(i)
        return out


@dataclass
class SeqHit:
    value: str
    basis: str


class DatasetBuilder:
    def build(self, paths: list[str | Path]) -> list[DataTable]:
        tables: list[DataTable] = []
        for path in paths:
            p = Path(path)
            suffix = p.suffix.lower()
            if suffix in (".xlsx", ".xlsm"):
                tables.extend(self._from_xlsx(p))
            elif suffix == ".docx":
                tables.extend(self._from_docx(p))
        return [t for t in tables if t.headers and t.rows]

    @staticmethod
    def _from_xlsx(path: Path) -> list[DataTable]:
        import openpyxl

        out: list[DataTable] = []
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if len(rows) < 2:
                continue
            headers = [str(v).strip() if v is not None else "" for v in rows[0]]
            body = [
                ["" if v is None else str(v).strip() for v in row]
                for row in rows[1:]
                if any(v is not None and str(v).strip() for v in row)
            ]
            out.append(DataTable(source=path.name, name=sheet.title, headers=headers, rows=body))
        wb.close()
        return out

    @staticmethod
    def _from_docx(path: Path) -> list[DataTable]:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(str(path))
        out: list[DataTable] = []
        caption = ""
        t_idx = 0
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                text = (Paragraph(child, doc).text or "").strip()
                if text and len(text) <= 40:
                    caption = text
            elif child.tag == qn("w:tbl"):
                t_idx += 1
                table = Table(child, doc)
                if len(table.rows) < 2:
                    continue
                headers = [c.text.strip() for c in table.rows[0].cells]
                body = [
                    [c.text.strip() for c in row.cells] for row in table.rows[1:]
                ]
                out.append(
                    DataTable(
                        source=path.name,
                        name=caption or f"表{t_idx}",
                        headers=headers,
                        rows=body,
                    )
                )
        return out


class SequenceResolver:
    """把 `活动1名称` 解析成结构化表格里的具体单元格。"""

    def __init__(self, tables: list[DataTable], attr_map: dict[str, str] | None = None) -> None:
        self.tables = tables
        # attr_map 来自模型：模板属性名 → 数据表真实列名（如 交付物 → Work Product）。
        # 只做补充，词表命中的优先，避免模型把明显的对应关系改坏。
        self.attr_map = attr_map or {}

    def resolve(self, placeholder: str) -> SeqHit | None:
        match = SEQ_RE.match(placeholder)
        if not match:
            return None
        group, seq, attr = match.group(1), int(match.group(2)), match.group(3)

        table = self.table_for(group)
        if table is None:
            return None

        col = table.col_index(attr)
        mapped = self.attr_map.get(attr, "")
        if col is None and mapped:
            col = table.index_of(mapped)
        if col is None:
            return None
        value = table.cell(seq - 1, col)
        if not value:
            return None

        header = table.headers[col]
        return SeqHit(
            value=value,
            basis=f"取自《{table.source}》「{table.name}」第 {seq} 行「{header}」列",
        )

    def unresolved(self, placeholders: list[str]) -> list[tuple[str, str, DataTable]]:
        """列出「表找到了、但属性名对不上任何列」的占位符，交给模型做语义映射。"""
        out: list[tuple[str, str, DataTable]] = []
        seen: set[str] = set()
        for ph in placeholders:
            match = SEQ_RE.match(ph or "")
            if not match:
                continue
            group, attr = match.group(1), match.group(3)
            if attr in seen or attr in self.attr_map:
                continue
            table = self.table_for(group)
            if table is None or table.col_index(attr) is not None:
                continue
            seen.add(attr)
            out.append((group, attr, table))
        return out

    def table_for(self, group: str) -> DataTable | None:
        """按集合名找最匹配的表：先看表名，再看表头。"""
        scored: list[tuple[int, DataTable]] = []
        for table in self.tables:
            score = 0
            if group in table.name:
                score += 3
            if group in table.header_blob:
                score += 1
            if score:
                scored.append((score, table))
        if not scored:
            return None
        scored.sort(key=lambda x: x[0], reverse=True)
        return scored[0][1]
