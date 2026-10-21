"""参考文件解析：把任意文档切成「带原文定位」的语料块。

定位信息（locator）是整个证据链的地基——用户点开一个字段，
必须能直接跳回「第 3 页第 2 段」看到原文，否则所谓溯源就是空话。
"""
from __future__ import annotations

import re
from pathlib import Path

from ..core.models import Chunk

CHUNK_SIZE = 600
CHUNK_OVERLAP = 80
SUPPORTED = {".docx", ".pdf", ".xlsx", ".xlsm", ".pptx", ".md", ".txt", ".markdown"}


def _split_long(text: str, size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    text = text.strip()
    if len(text) <= size:
        return [text] if text else []
    parts: list[str] = []
    step = size - overlap
    for start in range(0, len(text), step):
        piece = text[start:start + size].strip()
        if piece:
            parts.append(piece)
        if start + size >= len(text):
            break
    return parts


class DocumentParser:
    """按扩展名分派到具体解析器，统一产出 Chunk 列表。"""

    def parse(self, path: str | Path) -> list[Chunk]:
        path = Path(path)
        suffix = path.suffix.lower()
        if suffix not in SUPPORTED:
            raise ValueError(f"暂不支持的格式：{suffix}")

        handler = {
            ".docx": self._parse_docx,
            ".pdf": self._parse_pdf,
            ".xlsx": self._parse_xlsx,
            ".xlsm": self._parse_xlsx,
            ".pptx": self._parse_pptx,
            ".md": self._parse_text,
            ".markdown": self._parse_text,
            ".txt": self._parse_text,
        }[suffix]
        chunks = handler(path)
        return [c for c in chunks if c.text.strip()]

    # ---------- docx ----------

    def _parse_docx(self, path: Path) -> list[Chunk]:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        doc = Document(str(path))
        chunks: list[Chunk] = []
        p_idx = 0
        t_idx = 0

        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                p_idx += 1
                text = (Paragraph(child, doc).text or "").strip()
                for k, piece in enumerate(_split_long(text)):
                    chunks.append(
                        Chunk(
                            chunk_id=f"{path.stem}-p{p_idx}-{k}",
                            doc_name=path.name,
                            locator=f"第 {p_idx} 段",
                            text=piece,
                        )
                    )
            elif child.tag == qn("w:tbl"):
                t_idx += 1
                table = Table(child, doc)
                header = [c.text.strip() for c in table.rows[0].cells] if table.rows else []
                for r_i, row in enumerate(table.rows, start=1):
                    pairs = []
                    seen: set = set()
                    for c_i, cell in enumerate(row.cells):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        value = cell.text.strip()
                        if not value:
                            continue
                        key = header[c_i] if c_i < len(header) and header[c_i] else f"第{c_i + 1}列"
                        pairs.append(f"{key}：{value}")
                    if pairs:
                        chunks.append(
                            Chunk(
                                chunk_id=f"{path.stem}-t{t_idx}-r{r_i}",
                                doc_name=path.name,
                                locator=f"表 {t_idx} 第 {r_i} 行",
                                text="；".join(pairs),
                                section=header[0] if header else "",
                            )
                        )
        return chunks

    # ---------- pdf ----------

    def _parse_pdf(self, path: Path) -> list[Chunk]:
        import pdfplumber

        chunks: list[Chunk] = []
        with pdfplumber.open(str(path)) as pdf:
            for page_no, page in enumerate(pdf.pages, start=1):
                text = (page.extract_text() or "").strip()
                if not text:
                    continue
                for k, piece in enumerate(_split_long(text)):
                    chunks.append(
                        Chunk(
                            chunk_id=f"{path.stem}-pg{page_no}-{k}",
                            doc_name=path.name,
                            locator=f"第 {page_no} 页",
                            text=piece,
                        )
                    )
        return chunks

    # ---------- xlsx ----------

    def _parse_xlsx(self, path: Path) -> list[Chunk]:
        import openpyxl

        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        chunks: list[Chunk] = []
        for sheet in wb.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(v).strip() if v is not None else "" for v in rows[0]]
            for r_i, row in enumerate(rows[1:], start=2):
                pairs = [
                    f"{header[c_i] or f'第{c_i + 1}列'}：{value}"
                    for c_i, value in enumerate(row)
                    if value is not None and str(value).strip()
                ]
                if not pairs:
                    continue
                chunks.append(
                    Chunk(
                        chunk_id=f"{path.stem}-{sheet.title}-r{r_i}",
                        doc_name=path.name,
                        locator=f"{sheet.title}!第 {r_i} 行",
                        text="；".join(pairs),
                    )
                )
        wb.close()
        return chunks

    # ---------- pptx ----------

    def _parse_pptx(self, path: Path) -> list[Chunk]:
        from pptx import Presentation

        prs = Presentation(str(path))
        chunks: list[Chunk] = []
        for s_i, slide in enumerate(prs.slides, start=1):
            texts = [
                shape.text_frame.text.strip()
                for shape in slide.shapes
                if shape.has_text_frame and shape.text_frame.text.strip()
            ]
            if not texts:
                continue
            chunks.append(
                Chunk(
                    chunk_id=f"{path.stem}-s{s_i}",
                    doc_name=path.name,
                    locator=f"第 {s_i} 页幻灯片",
                    text="\n".join(texts),
                )
            )
        return chunks

    # ---------- 纯文本 ----------

    def _parse_text(self, path: Path) -> list[Chunk]:
        raw = path.read_text(encoding="utf-8", errors="ignore")
        chunks: list[Chunk] = []
        for p_idx, block in enumerate(re.split(r"\n\s*\n", raw), start=1):
            block = block.strip()
            if not block:
                continue
            for k, piece in enumerate(_split_long(block)):
                chunks.append(
                    Chunk(
                        chunk_id=f"{path.stem}-p{p_idx}-{k}",
                        doc_name=path.name,
                        locator=f"第 {p_idx} 段",
                        text=piece,
                    )
                )
        return chunks
