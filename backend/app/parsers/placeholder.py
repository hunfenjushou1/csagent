"""占位符模式模板解析。

相比「空白表格推断」，占位符模式是可判定的：模板作者已经把语义写在
`{项目名称}` 里了，不需要猜表头、猜哪个空是待填。这带来两个直接收益：

1. 字段名就是检索 query，召回和判定精度大幅提升；
2. 写回是精确替换，不会出现「填错格子」的问题。

兼容两种写法：`{{项目名称}}`（docxtpl 风格）和 `{项目名称}`。
同名占位符视为同一字段（全局同值），并记录出现次数——
模板里出现 5 次 `{姓名}` 说明作者想要同一个值，不是 5 个人。
需要多行数据时应在模板里写 `{姓名1}` `{姓名2}`。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from ..core.models import FieldKind, TemplateField, TemplateSpec, file_fingerprint
from .docx_template import _iter_blocks, infer_value_type

PH_RE = re.compile(r"\{\{\s*([^{}\n]{1,40}?)\s*\}\}|\{\s*([^{}\n]{1,40}?)\s*\}")
VALID_NAME = re.compile(r"^[\w\u4e00-\u9fa5 \-./]{1,40}$")


def count_placeholders(path: str | Path) -> int:
    """模式探测：占位符数量足够多时，说明这是一份占位符模板。"""
    doc = Document(str(path))
    total = 0
    for _, block in _iter_blocks(doc):
        if hasattr(block, "text"):
            total += _count_in(block.text or "")
        else:
            for row in block.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        total += _count_in(para.text)
    return total


def _count_in(text: str) -> int:
    return sum(1 for m in PH_RE.finditer(text) if VALID_NAME.match((m.group(1) or m.group(2) or "").strip()))


class PlaceholderAnalyzer:
    def analyze(self, path: str | Path, name: str = "") -> TemplateSpec:
        path = Path(path)
        doc = Document(str(path))

        collected: dict[str, TemplateField] = {}
        order: list[str] = []
        section = ""
        p_idx = -1
        t_idx = -1

        for kind, block in _iter_blocks(doc):
            if kind == "p":
                p_idx += 1
                text = (block.text or "").strip()
                if not text:
                    continue
                if self._is_heading(block, text):
                    section = text
                    continue
                self._collect(
                    text, section, f"第 {p_idx} 段",
                    paragraph_index=p_idx, collected=collected, order=order,
                )
            else:
                t_idx += 1
                for r_i, row in enumerate(block.rows):
                    seen: set = set()
                    for c_i, cell in enumerate(row.cells):
                        if cell._tc in seen:
                            continue
                        seen.add(cell._tc)
                        for para in cell.paragraphs:
                            self._collect(
                                para.text, section,
                                f"表 {t_idx + 1} 第 {r_i + 1} 行第 {c_i + 1} 列",
                                paragraph_index=None, collected=collected, order=order,
                                table_index=t_idx, row_index=r_i, col_index=c_i,
                            )

        # 页眉页脚里的占位符也要算（常见于「文件编号」「版本」）
        for sec in doc.sections:
            for label, hf in (("页眉", sec.header), ("页脚", sec.footer)):
                for para in hf.paragraphs:
                    self._collect(
                        para.text, section, label,
                        paragraph_index=None, collected=collected, order=order,
                    )

        fields = [collected[n] for n in order]
        for i, f in enumerate(fields, start=1):
            f.field_id = f"F{i:03d}"
            if f.occurrence > 1:
                f.hint += f"（该占位符在模板中出现 {f.occurrence} 次，将填写同一取值）"

        return TemplateSpec(
            template_id="",
            name=name or path.stem,
            source="local",
            source_ref=file_fingerprint(str(path)),
            fields=fields,
        )

    @staticmethod
    def _is_heading(para, text: str) -> bool:
        style = (para.style.name or "") if para.style is not None else ""
        return style.startswith(("Heading", "标题")) or (
            bool(re.match(r"^(\d+(\.\d+)*|[一二三四五六七八九十]+、)\s*\S", text)) and len(text) <= 40
        )

    def _collect(
        self,
        text: str,
        section: str,
        location: str,
        paragraph_index: int | None,
        collected: dict[str, TemplateField],
        order: list[str],
        table_index: int | None = None,
        row_index: int | None = None,
        col_index: int | None = None,
    ) -> None:
        for match in PH_RE.finditer(text or ""):
            raw = match.group(1) or match.group(2) or ""
            name = raw.strip()
            if not VALID_NAME.match(name):
                continue
            if name in collected:
                collected[name].occurrence += 1
                continue
            collected[name] = TemplateField(
                field_id="",
                kind=FieldKind.PLACEHOLDER,
                label=name,
                hint=f"占位符「{name}」，位于{location}，所属章节「{section or '未标注'}」",
                section=section,
                placeholder=name,
                value_type=infer_value_type(name),
                paragraph_index=paragraph_index,
                table_index=table_index,
                row_index=row_index,
                col_index=col_index,
                occurrence=1,
            )
            order.append(name)
