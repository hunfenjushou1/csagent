"""docx 模板结构分析器。

面对的是「空白表格 + 空白段落」这种没有任何占位符的模板，
所以待填位置只能靠结构推断：

- 表格：首行认作表头，数据行里的空白单元格即待填；合并单元格去重，
  「（填写说明）」「请填写」这类提示文字视作待填而非已填。
- 段落：形如「项目名称：」「XXX：____」的段落即待填；
  是否成立还要看下一段——若下一段是正文，说明这只是引导语，跳过。
- 复选框段落（□/☐）：识别为多选字段，交由用户勾选，不让模型猜。

推断必然有误差，所以每个字段都带上 location 与 hint，
让用户在审阅页能一眼看出「这个空我为什么要你填」。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..core.models import (
    FieldKind,
    TemplateField,
    TemplateSpec,
    ValueType,
    file_fingerprint,
)

PLACEHOLDER_RE = re.compile(r"^(.{2,40}?)[：:]\s*[_＿\-—]{2,}\s*$")
LABEL_ONLY_RE = re.compile(r"^(.{2,40}?)[：:]\s*$")
SECTION_NO_RE = re.compile(r"^(\d+(\.\d+)*|[一二三四五六七八九十]+、)\s*\S")
CHECKBOX_CHARS = ("□", "☐", "○", "◯")
NOTE_PREFIXES = ("注：", "注:", "说明：", "备注：", "（填写说明", "(填写说明")
INSTRUCTION_RE = re.compile(r"^[（(].{0,60}[)）]$|请填|此处填|填写说明|待填写")

LONG_TEXT_KEYS = ("描述", "说明", "内容", "措施", "要求", "方案", "概述", "活动",
                  "职责", "影响", "理由", "风险", "范围", "目标", "结论")
DATE_KEYS = ("时间", "日期", "期限", "起止")
NUMBER_KEYS = ("数量", "人数", "金额", "分值", "等级")
LEAD_IN_PREFIX = ("请", "依据", "如下", "见", "注", "按照", "根据", "为", "本表")


def _cell_text(cell) -> str:
    return "\n".join(p.text for p in cell.paragraphs).strip()


def _iter_blocks(doc):
    """按文档真实顺序交替产出段落与表格（doc.paragraphs 做不到这一点）。"""
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield "p", Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield "tbl", Table(child, doc)


def infer_value_type(label: str) -> ValueType:
    """字段名 → 值类型。占位符模式和结构推断模式共用。"""
    if any(k in label for k in DATE_KEYS):
        return ValueType.DATE
    if any(k in label for k in NUMBER_KEYS):
        return ValueType.NUMBER
    if any(k in label for k in LONG_TEXT_KEYS):
        return ValueType.LONG_TEXT
    if any(k in label for k in ("成员", "人员", "清单")):
        return ValueType.LIST
    return ValueType.TEXT


class DocxTemplateAnalyzer:
    def __init__(self, max_fields: int = 400) -> None:
        self.max_fields = max_fields

    def analyze(self, path: str | Path, name: str = "") -> TemplateSpec:
        path = Path(path)
        doc = Document(str(path))
        blocks = list(_iter_blocks(doc))

        fields: list[TemplateField] = []
        section = ""
        caption = ""
        p_idx = -1
        t_idx = -1
        seq = 0

        for i, (kind, block) in enumerate(blocks):
            if kind == "p":
                p_idx += 1
                text = (block.text or "").strip()
                if not text:
                    continue
                if self._is_heading(block, text):
                    section = text
                    caption = ""
                    continue
                # 引导语（「请从…中确认」「依据…如下」）不是表格标题，
                # 拿它当 caption 会让审阅页的定位信息变成废话
                if not self._is_lead_in(text):
                    caption = text
                field = self._paragraph_field(
                    text, p_idx, section, caption, blocks, i
                )
                if field is not None:
                    seq += 1
                    field.field_id = f"F{seq:03d}"
                    fields.append(field)
            else:
                t_idx += 1
                for field in self._table_fields(block, t_idx, section, caption):
                    seq += 1
                    field.field_id = f"F{seq:03d}"
                    fields.append(field)
                caption = ""

            if len(fields) >= self.max_fields:
                break

        return TemplateSpec(
            template_id="",
            name=name or path.stem,
            source="local",
            source_ref=file_fingerprint(str(path)),
            fields=fields,
        )

    # ---------- 段落 ----------

    def _is_heading(self, para: Paragraph, text: str) -> bool:
        style = (para.style.name or "") if para.style is not None else ""
        if style.startswith(("Heading", "标题")):
            return True
        return bool(SECTION_NO_RE.match(text)) and len(text) <= 40

    def _paragraph_field(
        self,
        text: str,
        p_idx: int,
        section: str,
        caption: str,
        blocks: list,
        pos: int,
    ) -> TemplateField | None:
        # 复选框段落：□选项A  □选项B
        if any(ch in text for ch in CHECKBOX_CHARS):
            options = self._checkbox_options(text)
            return self._field(
                kind=FieldKind.PARAGRAPH,
                label=f"勾选项：{options[0]}" if options else "勾选项",
                hint=f"章节「{section}」下的多选项，候选为：{'、'.join(options)}",
                section=section,
                caption=caption,
                value_type=ValueType.LIST,
                paragraph_index=p_idx,
            )

        match = PLACEHOLDER_RE.match(text)
        if match:
            label = match.group(1).strip()
            return self._field(
                kind=FieldKind.PARAGRAPH,
                label=label,
                hint=f"章节「{section}」中的字段「{label}」，用于填写对应内容",
                section=section,
                caption=caption,
                value_type=infer_value_type(label),
                paragraph_index=p_idx,
            )

        match = LABEL_ONLY_RE.match(text)
        if not match:
            return None
        label = match.group(1).strip()
        if len(label) < 2 or "。" in label:
            return None
        # 冒号后为空，但下一段是正文 → 这是引导语，不是待填字段
        nxt = self._next_text(blocks, pos)
        if nxt == "__table__":
            return None
        if nxt and len(nxt) > 30 and not any(ch in nxt for ch in CHECKBOX_CHARS):
            return None
        return self._field(
            kind=FieldKind.PARAGRAPH,
            label=label,
            hint=f"章节「{section}」中的字段「{label}」",
            section=section,
            caption=caption,
            value_type=infer_value_type(label),
            paragraph_index=p_idx,
        )

    @staticmethod
    def _next_text(blocks: list, pos: int) -> str:
        for kind, block in blocks[pos + 1:]:
            if kind == "p":
                text = (block.text or "").strip()
                if text:
                    return text
            else:
                return "__table__"
        return ""

    # ---------- 表格 ----------

    def _table_fields(
        self, table: Table, t_idx: int, section: str, caption: str
    ) -> list[TemplateField]:
        rows = table.rows
        if not rows:
            return []

        header = [_cell_text(c) for c in rows[0].cells]
        has_header = self._looks_like_header(header)
        columns = header if has_header else [f"第{i + 1}列" for i in range(len(header))]
        start = 1 if has_header else 0

        out: list[TemplateField] = []
        filled_cols: set[int] = set()   # 已生成过字段的列，用于合并整行留白
        merged_rows = 0

        for r_i in range(start, len(rows)):
            cells = rows[r_i].cells
            row_label = _cell_text(cells[0]) if cells else ""
            row_text = " ".join(_cell_text(c) for c in cells).strip()
            if any(row_text.startswith(p) for p in NOTE_PREFIXES):
                continue

            # 整行完全空白 = 模板预留的填写行，同一列只生成一个字段，
            # 否则 5 行留白会变成 5 组重复字段，把审阅页冲垮
            row_all_blank = not row_text
            if row_all_blank:
                merged_rows += 1

            seen: set = set()
            for c_i, cell in enumerate(cells):
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)

                text = _cell_text(cell)
                is_blank = not text
                is_hint = bool(text) and bool(INSTRUCTION_RE.search(text))
                if not is_blank and not is_hint:
                    continue
                if c_i >= len(columns):
                    continue
                if row_all_blank and c_i in filled_cols:
                    continue
                filled_cols.add(c_i)

                col = columns[c_i].strip() or f"第{c_i + 1}列"

                # 两列 key-value 表：「项目 | 内容」这种，语义标签在第一列
                if len(columns) == 2 and c_i == 1 and row_label:
                    label = row_label
                    hint = (
                        f"章节「{section}」表格「{caption or f'表{t_idx + 1}'}」中"
                        f"「{row_label}」对应的取值"
                    )
                else:
                    label = col if not row_label else f"{row_label} / {col}"
                    hint = (
                        f"章节「{section}」表格「{caption or f'表{t_idx + 1}'}」"
                        f"的「{col}」列"
                    )
                    if row_label:
                        hint += f"，本行对象为「{row_label}」"
                if is_hint:
                    hint += f"；模板提示：{text}"

                out.append(
                    self._field(
                        kind=FieldKind.TABLE_CELL,
                        label=label,
                        hint=hint,
                        section=section,
                        caption=caption,
                        value_type=infer_value_type(label),
                        table_index=t_idx,
                        row_index=r_i,
                        col_index=c_i,
                        current_value="" if is_hint else text,
                    )
                )

        if merged_rows > 1 and out:
            out[0].hint += f"（该表预留 {merged_rows} 行可批量填写，已合并为单条字段）"
        return out

    @staticmethod
    def _looks_like_header(texts: list[str]) -> bool:
        if len(texts) < 2:
            return False
        head = [t for t in texts if t]
        if len(head) < max(2, len(texts) // 2):
            return False
        return all(len(t) <= 24 for t in head)

    # ---------- 工具 ----------

    @staticmethod
    def _field(**kwargs) -> TemplateField:
        return TemplateField(field_id="", **kwargs)

    @staticmethod
    def _is_lead_in(text: str) -> bool:
        return text.startswith(LEAD_IN_PREFIX) or len(text) > 40

    @staticmethod
    def _checkbox_options(text: str) -> list[str]:
        pattern = "[" + "".join(CHECKBOX_CHARS) + "]"
        return [p.strip() for p in re.split(pattern, text) if p.strip()]

    @staticmethod
    def _infer_type(label: str) -> ValueType:
        return infer_value_type(label)
