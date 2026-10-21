"""语义对齐的离线测试：不需要 API Key，用假模型专门攻击校验层。

重点不是"模型能答对"，而是"模型答错的时候，正文会不会被污染"。
所以这里构造的全是恶意/错误输出：编造引文、编造取值、引用不存在的块、
以及最危险的一种——引文是真的、取值是另外编的。
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.core.models import Chunk, FieldKind, FieldStatus, TemplateField, ValueType  # noqa: E402
from app.services.aligner import SemanticAligner, is_traceable  # noqa: E402
from app.services.dataset import DataTable, SequenceResolver  # noqa: E402
from app.services.llm import ResultCache  # noqa: E402

PASSED: list[str] = []
FAILED: list[str] = []


def check(name: str, cond: bool, detail: str = "") -> None:
    if cond:
        PASSED.append(name)
        print(f"  [PASS] {name}")
    else:
        FAILED.append(name)
        print(f"  [FAIL] {name} {detail}")


class FakeClient:
    """假模型：按预设顺序吐结果，并记录每次调用。"""

    enabled = True

    class settings:  # noqa: N801
        model = "fake-deepseek"

    def __init__(self, responses: list[dict]) -> None:
        self.responses = list(responses)
        self.calls: list[str] = []
        self.tokens = 0

    def chat_json(self, system: str, user: str, max_tokens: int | None = None) -> dict:
        self.calls.append(user)
        if not self.responses:
            raise AssertionError("假模型被多调用了一次，说明缓存没生效")
        return self.responses.pop(0)


class DeadClient(FakeClient):
    def __init__(self) -> None:
        super().__init__([])

    def chat_json(self, system: str, user: str, max_tokens: int | None = None) -> dict:
        raise AssertionError("这一步不应产生模型调用（应全部命中缓存）")


def make_fields() -> list[TemplateField]:
    def f(fid: str, label: str, vtype: ValueType = ValueType.TEXT) -> TemplateField:
        return TemplateField(
            field_id=fid,
            kind=FieldKind.PLACEHOLDER,
            label=label,
            hint="章节「1」段落「" + label + "：{}」",
            placeholder=label,
            value_type=vtype,
        )

    return [
        f("F001", "网络安全负责人"),
        f("F002", "网络安全接口人"),
        f("F003", "交付物", ValueType.LIST),
        f("F004", "结束时间", ValueType.DATE),
        f("F005", "安全培训预算", ValueType.NUMBER),
    ]


def make_chunks() -> list[Chunk]:
    return [
        Chunk("a-1", "立项说明.docx", "第 3 段", "网络安全负责人：李娜；项目周期 2026 年 3 月至 9 月。"),
        Chunk(
            "b-2",
            "接口协议.xlsx",
            "网络安全活动!第 2 行",
            "序号：1；活动名称：网络安全需求分析；Work Product：需求规格说明书；Target Date：2026-05-30",
        ),
    ]


def test_traceable() -> None:
    print("\n[1] 取值可溯源性判定")
    check("日期换格式可溯源", is_traceable("2026-05-30", "Target Date：2026/5/30", ValueType.DATE))
    check("姓名照抄可溯源", is_traceable("李娜", "负责人：李娜", ValueType.TEXT))
    check("编造姓名被拒", not is_traceable("张伟", "负责人：李娜", ValueType.TEXT))
    check("编造日期被拒", not is_traceable("2026-08-01", "期限 2026-05-30", ValueType.DATE))
    check("短清单缺项被拒", not is_traceable("需求规格说明书、渗透测试报告", "Work Product：需求规格说明书", ValueType.LIST))
    check("编造金额被拒", not is_traceable("120000", "预算 80000 元", ValueType.NUMBER))


def test_verify_gate() -> None:
    print("\n[2] 校验闸门：模型说谎时的四道关")
    fields = make_fields()
    chunks = make_chunks()
    fake = FakeClient([
        {
            "results": [
                # 正确：引文真实、取值可溯源
                {"field_id": "F001", "value": "李娜", "evidence_chunk": "c1",
                 "quote": "网络安全负责人：李娜", "confidence": 0.9, "reason": "原文直接写明"},
                # 编造引文：资料里没有这句
                {"field_id": "F002", "value": "赵敏", "evidence_chunk": "c1",
                 "quote": "网络安全接口人：赵敏，负责对外沟通", "confidence": 0.9, "reason": "模型编的"},
                # 引文真实，但取值编造（最难识别的一种）
                {"field_id": "F003", "value": "需求规格说明书、渗透测试报告", "evidence_chunk": "c2",
                 "quote": "Work Product：需求规格说明书", "confidence": 0.8, "reason": "清单"},
                # 引用不存在的资料块
                {"field_id": "F004", "value": "2026-05-30", "evidence_chunk": "c99",
                 "quote": "Target Date：2026-05-30", "confidence": 0.9, "reason": "越界引用"},
                # 模型诚实地说找不到
                {"field_id": "F005", "value": None, "evidence_chunk": None, "quote": None,
                 "confidence": 0.0, "reason": "资料中只有人力安排，未出现任何培训预算金额"},
            ]
        }
    ])
    aligner = SemanticAligner(fake, cache=ResultCache(None))
    stats = aligner.align(fields, chunks, question="按接口协议填写网络安全计划")

    by_id = {f.field_id: f for f in fields}

    f1 = by_id["F001"]
    check("真实依据被采纳", f1.status == FieldStatus.PENDING and f1.value == "李娜", f"{f1.status} {f1.value}")
    check("采纳结果挂上了 chunk 证据", f1.evidence == ["a-1"], str(f1.evidence))
    check("采纳结果标注需人工核对", "请核对" in f1.basis, f1.basis)

    f2 = by_id["F002"]
    check("编造引文被拒绝", f2.status == FieldStatus.MISSING and f2.value == "", f"{f2.status} {f2.value}")
    check("拒绝理由说明是编造", "找不到" in f2.basis or "编造" in f2.basis, f2.basis)
    check("拒绝后仍向用户追问", bool(f2.question) and "F002" not in f2.question, f2.question)

    f3 = by_id["F003"]
    check("引文真但取值假被拒绝", f3.status == FieldStatus.MISSING and "渗透测试报告" not in f3.value, f3.value)

    f4 = by_id["F004"]
    check("引用不存在的块被拒绝", f4.status == FieldStatus.MISSING, f"{f4.status} {f4.value}")

    f5 = by_id["F005"]
    check("模型找不到时保留追问", f5.status == FieldStatus.MISSING and "培训预算" in f5.question, f5.question)
    check("追问里带上模型的复核说明", "培训预算金额" in f5.question, f5.question)

    check("统计口径正确", stats.accepted == 1 and stats.rejected == 3 and stats.empty == 1,
          f"accepted={stats.accepted} rejected={stats.rejected} empty={stats.empty}")

    print("\n[3] 缓存：同一份资料不重复调模型")
    fields2 = make_fields()
    aligner2 = SemanticAligner(DeadClient(), cache=aligner.cache)
    stats2 = aligner2.align(fields2, chunks)
    check("第二次全部命中缓存", stats2.cached == 5 and stats2.attempted == 5, f"cached={stats2.cached}")
    check("缓存还原了采纳结果", fields2[0].status == FieldStatus.PENDING and fields2[0].value == "李娜")

    print("\n[4] 没有 Key 时的降级")
    off = SemanticAligner(None)
    stats3 = off.align(make_fields(), chunks)
    check("未配置 Key 不报错", stats3.enabled is False and "规则模式" in stats3.note, stats3.note)


def test_column_mapping() -> None:
    print("\n[5] 列名语义映射（{活动1交付物} → 验收材料）")
    table = DataTable(
        source="接口协议.xlsx",
        name="网络安全活动",
        headers=["序号", "活动名称", "验收材料", "Target Date"],
        rows=[["1", "网络安全需求分析", "需求规格说明书", "2026-05-30"]],
    )
    field_obj = TemplateField(
        field_id="F010",
        kind=FieldKind.PLACEHOLDER,
        label="交付物",
        hint="",
        placeholder="活动1交付物",
    )

    plain = SequenceResolver([table])
    check("未映射前词表对不上", plain.resolve("活动1交付物") is None)

    fake = FakeClient([{"mapping": [{"attr": "交付物", "column": "验收材料"}]}])
    aligner = SemanticAligner(fake, cache=ResultCache(None))
    mapping, note = aligner.align_columns([field_obj], [table])
    check("模型给出了列映射", mapping == {"交付物": "验收材料"}, str(mapping))

    resolver = SequenceResolver([table], attr_map=mapping)
    hit = resolver.resolve("活动1交付物")
    check("映射后精确落到单元格", hit is not None and hit.value == "需求规格说明书", str(hit))

    fake2 = FakeClient([{"mapping": [{"attr": "交付物", "column": "不存在的列"}]}])
    bad, _ = SemanticAligner(fake2, cache=ResultCache(None)).align_columns([field_obj], [table])
    check("模型编造列名被拒绝", bad == {}, str(bad))


def test_wiring() -> None:
    """端到端接线：假模型 + 真链路，确认 _run_pipeline 确实把语义对齐接进去了。"""
    print("\n[6] 端到端接线（假模型跑真实链路）")
    import re
    import tempfile
    from pathlib import Path

    from docx import Document

    import app.main as main_mod
    from app.services.llm import parse_json_object

    check("解析带 ``` 包裹的 JSON", parse_json_object('```json\n{"a":1}\n```')["a"] == 1)
    check("解析带前后解释的 JSON", parse_json_object('结果如下 {"a":2} 完')["a"] == 2)

    class SmartFake(FakeClient):
        def chat_json(self, system: str, user: str, max_tokens: int | None = None) -> dict:
            self.calls.append(user)
            if "列名" in system:
                return {"mapping": []}
            out = []
            for line in user.splitlines():
                m = re.match(r"(F\d+) \| 字段名：(.+?) \|", line)
                if not m:
                    continue
                fid, label = m.group(1), m.group(2)
                if label == "活动负责人":
                    out.append({
                        "field_id": fid, "value": "李娜", "evidence_chunk": "c1",
                        "quote": "本项目的对接人为李娜（整车电子部）",
                        "confidence": 0.85, "reason": "资料里的对接人即活动负责人",
                    })
                elif label == "安全预算":
                    out.append({
                        "field_id": fid, "value": None, "evidence_chunk": None, "quote": None,
                        "confidence": 0.0, "reason": "资料中没有任何预算金额",
                    })
            return {"results": out}

    tmp = Path(tempfile.mkdtemp())
    doc = Document()
    doc.add_heading("网络安全工作说明", level=0)
    doc.add_paragraph("活动负责人：{活动负责人}")
    doc.add_paragraph("安全预算：{安全预算}")
    doc.add_paragraph("项目名称：{项目名称}")
    tpl = tmp / "tpl.docx"
    doc.save(str(tpl))
    ref = tmp / "说明.txt"
    ref.write_text(
        "本项目的对接人为李娜（整车电子部），联系电话 13800000001。\n\n"
        "项目按计划推进，各阶段活动由供应商与整车电子部联合执行。",
        encoding="utf-8",
    )

    original = main_mod._make_aligner
    main_mod._make_aligner = lambda: SemanticAligner(SmartFake([]), cache=ResultCache(None))
    try:
        payload = main_mod._run_pipeline("testwiring", tpl, [ref], "按说明填写")
    finally:
        main_mod._make_aligner = original
        import shutil
        shutil.rmtree(main_mod.TASKS / "testwiring", ignore_errors=True)
        shutil.rmtree(tmp, ignore_errors=True)

    fields = {f["label"]: f for f in payload["fields"]}
    check("payload 带 llm 段", payload.get("llm", {}).get("enabled") is True, str(payload.get("llm")))
    check("语义对齐采纳 1 项", payload["llm"]["accepted"] == 1, str(payload["llm"]))

    owner = fields.get("活动负责人", {})
    check("规则找不到的字段被模型补上", owner.get("status") == "pending_confirm", str(owner.get("status")))
    check("补上的值正确", owner.get("value") == "李娜", str(owner.get("value")))
    check("依据带原文定位", "第 1 段" in (owner.get("basis") or ""), owner.get("basis"))
    check("证据链指向真实语料块", bool(owner.get("evidence")), str(owner.get("evidence")))

    budget = fields.get("安全预算", {})
    check("模型说没有的字段仍为缺失", budget.get("status") == "missing", str(budget.get("status")))
    check("追问里说明模型也没找到", "预算金额" in (budget.get("question") or ""), budget.get("question"))


def test_http_client() -> None:
    """HTTP 层：用本地假 DeepSeek 服务验证请求格式与解析，不花一分钱。"""
    print("\n[7] HTTP 客户端（本地假服务）")
    import json as _json
    import threading
    from http.server import BaseHTTPRequestHandler, HTTPServer

    from app.services.llm import DeepSeekClient, LLMSettings

    seen: dict = {}

    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):  # 静音
            return

        def do_POST(self):
            length = int(self.headers.get("content-length") or 0)
            body = _json.loads(self.rfile.read(length) or b"{}")
            seen["path"] = self.path
            seen["auth"] = self.headers.get("authorization")
            seen["format"] = (body.get("response_format") or {}).get("type")
            seen["model"] = body.get("model")
            content = _json.dumps({"ok": True}, ensure_ascii=False)
            payload = _json.dumps(
                {"choices": [{"message": {"content": content}}], "usage": {"total_tokens": 42}}
            ).encode()
            self.send_response(200)
            self.send_header("content-type", "application/json")
            self.send_header("content-length", str(len(payload)))
            self.end_headers()
            self.wfile.write(payload)

    server = HTTPServer(("127.0.0.1", 0), Handler)
    port = server.server_address[1]
    threading.Thread(target=server.serve_forever, daemon=True).start()

    try:
        settings = LLMSettings(
            api_key="sk-test", base_url=f"http://127.0.0.1:{port}", model="deepseek-chat"
        )
        client = DeepSeekClient(settings)
        result = client.ping()
        check("连通性测试返回成功", result.get("ok") is True, str(result))
        check("请求打到 /chat/completions", seen.get("path") == "/chat/completions", str(seen.get("path")))
        check("带上 Bearer 鉴权", seen.get("auth") == "Bearer sk-test", str(seen.get("auth")))
        check("启用 JSON 输出模式", seen.get("format") == "json_object", str(seen.get("format")))
        check("token 计数被记录", client.tokens == 42, str(client.tokens))
    finally:
        server.shutdown()

    class Failing(BaseHTTPRequestHandler):
        def log_message(self, *args):
            return

        def do_POST(self):
            body = b'{"error":{"message":"invalid api key"}}'
            self.send_response(401)
            self.send_header("content-length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

    server2 = HTTPServer(("127.0.0.1", 0), Failing)
    threading.Thread(target=server2.serve_forever, daemon=True).start()
    try:
        bad = DeepSeekClient(
            LLMSettings(api_key="sk-bad", base_url=f"http://127.0.0.1:{server2.server_address[1]}")
        )
        bad_result = bad.ping()
        check(
            "Key 错误时不重试、返回原因",
            bad_result.get("ok") is False and "401" in str(bad_result.get("error")) and bad.calls == 1,
            f"{bad_result} calls={bad.calls}",
        )
    finally:
        server2.shutdown()


def main() -> int:
    test_traceable()
    test_verify_gate()
    test_column_mapping()
    test_wiring()
    test_http_client()
    print(f"\n通过 {len(PASSED)} 项，失败 {len(FAILED)} 项")
    if FAILED:
        for name in FAILED:
            print(f"  ! {name}")
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
