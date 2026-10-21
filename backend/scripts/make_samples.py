"""生成占位符模式的模板与参考文件。

对齐 plan.md 里「网络安全计划」的要求：
- 模板用 {占位符} 标注待填位置，不再靠结构推断
- 成员清单（xlsx）→ 提炼成员 → 逐人占位符 {成员N姓名}
- 接口协议（xlsx）→ 提炼活动，关注 Work Product 字段
- 活动起止时间可由 Target Date 类字段预填充
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
TPL_DIR = ROOT / "storage" / "templates"
UP_DIR = ROOT / "storage" / "uploads"


def _table(doc, header: list[str], data_rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(data_rows), cols=len(header))
    table.style = "Table Grid"
    for i, name in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = name
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for r_i, row_values in enumerate(data_rows, start=1):
        for c_i, value in enumerate(row_values):
            table.rows[r_i].cells[c_i].text = value
    return table


def build_template() -> Path:
    """占位符版网络安全计划模板。"""
    TPL_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading("{项目名称} 网络安全计划", level=0)
    doc.add_paragraph("文件编号：{文件编号}　　版本：{版本号}　　编制日期：{编制日期}")

    doc.add_heading("1 项目基本信息", level=1)
    _table(
        doc,
        ["项目", "内容"],
        [
            ["项目名称", "{项目名称}"],
            ["项目编号", "{项目编号}"],
            ["客户名称", "{客户名称}"],
            ["项目周期", "{项目周期}"],
            ["产品类型", "{产品类型}"],
            ["项目经理", "{项目经理}"],
        ],
    )

    doc.add_heading("2 网络安全相关人员", level=1)
    doc.add_paragraph("经项目组确认，本项目网络安全相关人员及其职责如下：")
    # 故意留 6 行而成员只有 4 人：验证「数据比模板行少」时能自动缩行
    _table(
        doc,
        ["序号", "姓名", "部门", "网络安全角色", "联系方式"],
        [
            ["1", "{成员1姓名}", "{成员1部门}", "{成员1角色}", "{成员1联系方式}"],
            ["2", "{成员2姓名}", "{成员2部门}", "{成员2角色}", "{成员2联系方式}"],
            ["3", "{成员3姓名}", "{成员3部门}", "{成员3角色}", "{成员3联系方式}"],
            ["4", "{成员4姓名}", "{成员4部门}", "{成员4角色}", "{成员4联系方式}"],
            ["5", "{成员5姓名}", "{成员5部门}", "{成员5角色}", "{成员5联系方式}"],
            ["6", "{成员6姓名}", "{成员6部门}", "{成员6角色}", "{成员6联系方式}"],
        ],
    )
    doc.add_paragraph("网络安全负责人：{网络安全负责人}")
    doc.add_paragraph("网络安全接口人：{网络安全接口人}")

    doc.add_heading("3 网络安全活动", level=1)
    doc.add_paragraph("依据网络安全接口协议，本项目需执行的网络安全活动及时间安排如下：")
    # 故意只写 4 行而协议有 5 条活动：验证「数据比模板行多」时能自动扩行
    _table(
        doc,
        ["活动编号", "活动名称", "责任人", "开始时间", "结束时间", "交付物"],
        [
            [
                "{活动1编号}", "{活动1名称}", "{活动1责任人}",
                "{活动1开始时间}", "{活动1结束时间}", "{活动1交付物}",
            ],
            [
                "{活动2编号}", "{活动2名称}", "{活动2责任人}",
                "{活动2开始时间}", "{活动2结束时间}", "{活动2交付物}",
            ],
            [
                "{活动3编号}", "{活动3名称}", "{活动3责任人}",
                "{活动3开始时间}", "{活动3结束时间}", "{活动3交付物}",
            ],
            [
                "{活动4编号}", "{活动4名称}", "{活动4责任人}",
                "{活动4开始时间}", "{活动4结束时间}", "{活动4交付物}",
            ],
        ],
    )

    doc.add_heading("4 裁剪说明", level=1)
    doc.add_paragraph("裁剪活动：{裁剪活动}")
    doc.add_paragraph("裁剪理由：{裁剪理由}")

    doc.add_heading("5 其他", level=1)
    doc.add_paragraph("选项：{供应商参与}")

    path = TPL_DIR / "网络安全计划模板.docx"
    doc.save(str(path))
    return path


def build_project_charter() -> Path:
    UP_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("电池控制器 BMS-2026 项目立项说明", level=0)

    doc.add_heading("一、项目概况", level=1)
    doc.add_paragraph(
        "项目名称：电池控制器 BMS-2026 平台化开发项目。"
        "项目编号：PRJ-2026-0317。客户名称：华东新能源汽车股份有限公司。"
    )
    doc.add_paragraph(
        "项目周期：2026 年 3 月至 2027 年 6 月。"
        "产品类型：动力电池控制器。项目经理：张伟。"
    )
    doc.add_paragraph("文件编号：CS-PLAN-2026-0317。版本号：V1.0。编制日期：2026-03-05。")

    doc.add_heading("二、网络安全范围", level=1)
    doc.add_paragraph(
        "本项目网络安全域活动覆盖 BMS 控制器本体、CAN 通信接口与 OTA 升级通道。"
        "需在概念阶段完成 TARA 分析，在开发阶段输出网络安全概念与规范。"
    )

    path = UP_DIR / "项目立项说明.docx"
    doc.save(str(path))
    return path


def build_member_list() -> Path:
    UP_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "成员清单"
    ws.append(["姓名", "部门", "岗位", "邮箱", "联系方式"])
    ws.append(["李娜", "整车电子部", "网络安全工程师", "lina@example.com", "13800000001"])
    ws.append(["王强", "底盘控制部", "软件工程师", "wangqiang@example.com", "13800000002"])
    ws.append(["陈静", "质量部", "质量工程师", "chenjing@example.com", "13800000003"])
    ws.append(["刘洋", "采购部", "供应商管理", "liuyang@example.com", "13800000004"])
    path = UP_DIR / "项目组成员清单.xlsx"
    wb.save(str(path))
    return path


def build_interface_agreement() -> Path:
    """接口协议改为 excel，带 Work Product 与 Target Date，供活动提炼与时间预填。"""
    UP_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "网络安全活动"
    ws.append(
        ["序号", "活动编号", "活动名称", "Work Product（工作产物）", "责任方", "Target Date"]
    )
    rows = [
        [1, "CS-A01", "网络安全需求分析", "网络安全需求规格说明书", "供应商", "2026-05-30"],
        [2, "CS-A02", "威胁分析与风险评估（TARA）", "TARA 报告", "联合", "2026-07-15"],
        [3, "CS-A03", "网络安全概念设计", "网络安全概念文档", "供应商", "2026-09-30"],
        [4, "CS-A04", "渗透测试", "渗透测试报告", "第三方", "2026-12-20"],
        [5, "CS-A05", "网络安全验证与确认", "验证报告", "供应商", "2027-03-15"],
    ]
    for row in rows:
        ws.append(row)

    ws2 = wb.create_sheet("接口职责")
    ws2.append(["事项", "整车厂", "供应商"])
    ws2.append(["TARA 输入", "提供整车级威胁场景", "参与评估"])
    ws2.append(["控制器本体设计", "评审", "负责"])

    path = UP_DIR / "网络安全接口协议.xlsx"
    wb.save(str(path))
    return path


def build_legacy_template() -> Path:
    """保留一份无占位符模板，用于回归结构推断模式。"""
    legacy_dir = TPL_DIR / "legacy"
    legacy_dir.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("XX 项目网络安全计划（结构推断示例）", level=0)
    doc.add_heading("1 项目基本信息", level=1)
    _table(doc, ["项目", "内容"], [["项目名称", ""], ["项目编号", ""]])
    doc.add_heading("2 网络安全相关人员", level=1)
    _table(doc, ["序号", "姓名", "部门"], [["", "", ""], ["", "", ""]])
    doc.add_paragraph("网络安全负责人：")
    path = legacy_dir / "结构推断模板.docx"
    doc.save(str(path))
    return path


def main() -> None:
    for fn in (
        build_template,
        build_project_charter,
        build_member_list,
        build_interface_agreement,
        build_legacy_template,
    ):
        print(f"[ok] {fn()}")


if __name__ == "__main__":
    main()
