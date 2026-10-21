"""端到端冒烟：模板拆解 → 参考文件解析 → 逐字段检索 → 三态判定。

跑法（在 backend 目录下）：
    python scripts/smoke.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.core.models import FieldStatus  # noqa: E402
from app.parsers.doc_parser import DocumentParser  # noqa: E402
from app.parsers.analyzer import analyze_template  # noqa: E402
from app.services.dataset import DatasetBuilder  # noqa: E402
from app.services.expander import expand_fields  # noqa: E402
from app.services.exporter import DocxExporter  # noqa: E402
from app.services.retriever import FieldRetriever  # noqa: E402

TPL = BACKEND / "storage" / "templates" / "网络安全计划模板.docx"
UPLOADS = [
    BACKEND / "storage" / "uploads" / "项目立项说明.docx",
    BACKEND / "storage" / "uploads" / "项目组成员清单.xlsx",
    BACKEND / "storage" / "uploads" / "网络安全接口协议.xlsx",
]
OUT = BACKEND / "storage" / "outputs" / "smoke_report.json"


def main() -> int:
    if not TPL.exists():
        print(f"[fail] 模板不存在：{TPL}\n先运行 python scripts/make_samples.py")
        return 1

    spec = analyze_template(TPL)
    print(f"模式：{spec.source}，指纹 {spec.source_ref}")
    print(f"\n=== 模板拆解：{spec.name}（指纹 {spec.source_ref}）===")
    print(f"共识别 {len(spec.fields)} 个字段，其中空白 {len(spec.blank_fields)} 个\n")
    for f in spec.fields:
        if f.kind.value == "placeholder":
            if f.table_index is not None:
                loc = f"表{f.table_index + 1} r{f.row_index}c{f.col_index}"
            elif f.paragraph_index is not None:
                loc = f"段{f.paragraph_index}"
            else:
                loc = "页眉页脚"
        elif f.kind.value == "table_cell":
            loc = f"表{f.table_index + 1} r{f.row_index}c{f.col_index}"
        else:
            loc = f"段{f.paragraph_index}"
        print(f"  {f.field_id} [{f.kind.value:11s}] {f.label[:30]:32s} {loc}")

    parser = DocumentParser()
    chunks = []
    for path in UPLOADS:
        if not path.exists():
            print(f"[warn] 缺少参考文件：{path.name}")
            continue
        got = parser.parse(path)
        chunks.extend(got)
        print(f"\n[parse] {path.name} → {len(got)} 块")
    print(f"\n语料合计 {len(chunks)} 块")

    tables = DatasetBuilder().build(UPLOADS)
    print(f"结构化表格 {len(tables)} 张：" + "、".join(f"{t.source}/{t.name}({len(t.rows)}行)" for t in tables))

    before = len(spec.fields)
    spec.fields, plan = expand_fields(spec.fields, tables)
    print("\n=== 动态扩行 ===")
    if not plan:
        print("  （无集合型占位符，或数据源未命中，保持模板行数）")
    for change in plan:
        print(
            f"  {change.collection.name}：模板 {change.collection.max_seq} 行"
            f" → 数据 {change.target_rows} 行（{change.source}）"
            f"  新增序号 {change.added or '—'}  裁掉序号 {change.removed or '—'}"
        )
    print(f"  字段数 {before} → {len(spec.fields)}")

    retriever = FieldRetriever(chunks, tables=tables)
    retriever.fill_fields(spec.fields)

    stats = {}
    for f in spec.fields:
        stats[f.status.value] = stats.get(f.status.value, 0) + 1

    print("\n=== 判定结果 ===")
    for f in spec.fields:
        mark = {"filled": "PASS", "pending_confirm": "CAND", "missing": "MISS"}[f.status.value]
        ev = f.evidence[0] if f.evidence else "-"
        print(f"  [{mark}] {f.label[:34]:36s} conf={f.confidence:.2f}  {f.basis[:40]}")

    print(f"\n统计：{json.dumps(stats, ensure_ascii=False)}")
    miss = [f for f in spec.fields if f.status == FieldStatus.MISSING]
    if miss:
        print(f"\n=== 需向用户追问 {len(miss)} 项 ===")
        for f in miss[:6]:
            print(f"  · {f.question}")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    report = {
        "template": spec.to_dict(),
        "chunk_count": len(chunks),
        "stats": stats,
        "fields": [f.to_dict() for f in spec.fields],
        "questions": [f.question for f in miss],
        "expansion": [c.to_dict() for c in plan],
    }
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n报告已写入 {OUT}")

    verify_export(spec.fields)
    return 0


def verify_export(fields) -> None:
    """导出并核验：表格行数是否跟数据条数一致、有没有残留裸占位符。"""
    from docx import Document

    from app.parsers.placeholder import PH_RE, VALID_NAME

    out = BACKEND / "storage" / "outputs" / "smoke导出.docx"
    DocxExporter().render(TPL, [f.to_dict() for f in fields], out, task_id="smoke")

    doc = Document(str(out))
    print("\n=== 导出核验 ===")
    for i, table in enumerate(doc.tables[:3]):
        head = " / ".join(c.text.strip() for c in table.rows[0].cells)[:44]
        print(f"  表{i + 1} {len(table.rows) - 1} 行数据  [{head}]")
        if i > 0:
            for row in table.rows[1:]:
                print("     ", [c.text.strip()[:16] for c in row.cells])

    leftovers = []
    for p in doc.paragraphs:
        for m in PH_RE.finditer(p.text or ""):
            name = (m.group(1) or m.group(2) or "").strip()
            if VALID_NAME.match(name):
                leftovers.append(name)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for m in PH_RE.finditer(p.text or ""):
                        name = (m.group(1) or m.group(2) or "").strip()
                        if VALID_NAME.match(name):
                            leftovers.append(name)

    if leftovers:
        print(f"  [FAIL] 残留裸占位符 {len(set(leftovers))} 个：{sorted(set(leftovers))[:8]}")
    else:
        print("  [PASS] 无残留占位符")


if __name__ == "__main__":
    raise SystemExit(main())
